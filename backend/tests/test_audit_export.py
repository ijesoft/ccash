import io
import uuid


def _sample_row():
    return {
        "id": str(uuid.uuid4()),
        "actor_email": "alice@ccash.ph",
        "action": "user.login",
        "resource_type": "user",
        "resource_id": "user-123-ref",
        "summary": "Alice logged in",
        "created_at": "2026-10-09T00:00:00+00:00",
    }


def test_audit_log_excel_has_grid_columns():
    from openpyxl import load_workbook

    from app.domains.admin.reports import generate_audit_log_excel

    row = _sample_row()
    data = generate_audit_log_excel([row])

    assert data[:2] == b"PK"

    wb = load_workbook(filename=io.BytesIO(data))
    ws = wb.active
    flat = [str(c.value) for r in ws.iter_rows() for c in r if c.value is not None]

    for header in ["Time", "Actor", "Action", "Resource", "Details"]:
        assert header in flat
    assert row["actor_email"] in flat
    assert row["resource_id"] not in flat


def test_audit_log_excel_handles_empty():
    from openpyxl import load_workbook

    from app.domains.admin.reports import generate_audit_log_excel

    data = generate_audit_log_excel([])

    assert data[:2] == b"PK"

    wb = load_workbook(filename=io.BytesIO(data))
    ws = wb.active
    flat = [str(c.value) for r in ws.iter_rows() for c in r if c.value is not None]

    for header in ["Time", "Actor", "Action", "Resource", "Details"]:
        assert header in flat


def test_audit_export_endpoint_requires_perm():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    import app.api.reports as reports_api
    from app.core.security import create_access_token

    test_app = FastAPI()
    test_app.include_router(reports_api.admin_router, prefix="/admin/reports")
    client = TestClient(test_app, raise_server_exceptions=False)

    assert client.get("/admin/reports/audit-log.xlsx").status_code == 401

    limited_token = create_access_token(str(uuid.uuid4()), scopes=["wallet:read", "audit:read"])
    res = client.get(
        "/admin/reports/audit-log.xlsx",
        headers={"Authorization": f"Bearer {limited_token}"},
    )
    assert res.status_code == 403
