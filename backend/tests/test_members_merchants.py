"""Members/Merchants sign-up fields, admin add/batch-upload, and report export.

Pure-function pieces (merchant ID generation, batch-file parsing, report byte
generation) need no DB and run everywhere. The admin create_member tests
follow the existing test_rbac.py convention: real Postgres via the session
fixture, service layer exercised directly, no HTTP.
"""

import io

import pytest

from app.core.errors import ValidationError
from app.domains.admin.batch_import import BatchImportError, parse_member_rows
from app.domains.admin.reports import generate_all_transactions_excel, generate_all_transactions_pdf
from app.domains.merchants.policy import (
    MERCHANT_ID_DIGITS,
    generate_merchant_id_no,
    is_valid_merchant_id_no,
)
from app.domains.transactions.reports import generate_transaction_history_excel, generate_transaction_history_pdf

# --------------------------------------------------------------- merchant ID


def test_generate_merchant_id_no_shape():
    for _ in range(20):
        value = generate_merchant_id_no()
        assert value.startswith("M")
        assert len(value) == 1 + MERCHANT_ID_DIGITS
        assert value[1:].isdigit()
        assert is_valid_merchant_id_no(value)


@pytest.mark.parametrize(
    "value",
    ["123456789", "M12345678", "M1234567890", "Mabcdefghi", "m123456789", ""],
)
def test_is_valid_merchant_id_no_rejects_malformed(value):
    assert not is_valid_merchant_id_no(value)


# -------------------------------------------------------------- batch import


CSV_HEADER = "ID No.,Last Name,First Name,Middle Name,Mobile,Email\n"


def test_parse_member_rows_csv_happy_path():
    content = (CSV_HEADER + "123456789,Dela Cruz,Juan,Santos,09171234567,juan@example.com\n").encode()
    rows = parse_member_rows(content, "members.csv")
    assert rows == [
        {
            "id_no": "123456789",
            "last_name": "Dela Cruz",
            "first_name": "Juan",
            "middle_name": "Santos",
            "mobile": "09171234567",
            "email": "juan@example.com",
        }
    ]


def test_parse_member_rows_skips_blank_rows():
    content = (CSV_HEADER + "\n123456789,Dela Cruz,Juan,,09171234567,juan@example.com\n,,,,,\n").encode()
    rows = parse_member_rows(content, "members.csv")
    assert len(rows) == 1


def test_parse_member_rows_missing_required_column_raises():
    content = b"Last Name,First Name,Mobile,Email\nDela Cruz,Juan,09171234567,juan@example.com\n"
    with pytest.raises(BatchImportError):
        parse_member_rows(content, "members.csv")


def test_parse_member_rows_rejects_unsupported_extension():
    with pytest.raises(BatchImportError):
        parse_member_rows(b"whatever", "members.txt")


def test_parse_member_rows_xlsx_happy_path():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["id_no", "last_name", "first_name", "middle_name", "mobile", "email"])
    # Excel often stores a numeric-looking cell as a number; id_no must survive as text.
    ws.append([123456789, "Dela Cruz", "Juan", "Santos", "09171234567", "juan@example.com"])
    buffer = io.BytesIO()
    wb.save(buffer)

    rows = parse_member_rows(buffer.getvalue(), "members.xlsx")
    assert rows[0]["id_no"] == "123456789"
    assert rows[0]["mobile"] == "09171234567"


# -------------------------------------------------------------- report bytes


def test_transaction_history_pdf_and_excel_handle_empty_history():
    pdf_bytes = generate_transaction_history_pdf("Alice Doe", [("ID No.", "000000002")], [])
    assert pdf_bytes.startswith(b"%PDF")

    xlsx_bytes = generate_transaction_history_excel("Alice Doe", [("ID No.", "000000002")], [])
    assert xlsx_bytes.startswith(b"PK")  # xlsx is a zip archive


def test_admin_all_transactions_pdf_and_excel_handle_empty_ledger():
    pdf_bytes = generate_all_transactions_pdf([])
    assert pdf_bytes.startswith(b"%PDF")

    xlsx_bytes = generate_all_transactions_excel([])
    assert xlsx_bytes.startswith(b"PK")


def test_admin_all_transactions_excel_includes_row_data():
    rows = [
        {
            "created_at": "2026-09-18 10:00",
            "reference": "CC260918ABCD1234",
            "type": "SEND",
            "status": "SUCCESS",
            "from": "Alice Doe",
            "to": "Sample Sari-Sari Store",
            "amount_cents": 10000,
            "fee_cents": 0,
            "net_amount_cents": 10000,
            "description": "Payment",
        }
    ]
    xlsx_bytes = generate_all_transactions_excel(rows)
    assert xlsx_bytes.startswith(b"PK")

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    values = [cell.value for row in ws.iter_rows() for cell in row]
    assert "Sample Sari-Sari Store" in values
    assert "CC260918ABCD1234" in values


# ------------------------------------------------------- admin create_member


async def test_admin_create_member_creates_active_verified_user_with_wallet(session):
    from app.domains.admin.service import AdminService
    from app.domains.auth.models import UserRole, UserStatus
    from app.domains.wallets.repository import WalletRepository

    service = AdminService(session)
    user, temp_password = await service.create_member(
        id_no="000000099",
        first_name="Juan",
        last_name="Dela Cruz",
        email="juan.batch@ccash.test",
        phone="09991234567",
        middle_name="Santos",
    )

    assert user.role == UserRole.MEMBER
    assert user.status == UserStatus.ACTIVE
    assert user.is_verified is True
    assert user.id_no == "000000099"
    assert len(temp_password) == 12

    wallet = await WalletRepository(session).get_by_user_id(user.id)
    assert wallet is not None
    assert wallet.balance_cents == 0


async def test_admin_create_member_rejects_duplicate_id_no(session):
    from app.domains.admin.service import AdminService

    service = AdminService(session)
    await service.create_member(
        id_no="000000100",
        first_name="Juan",
        last_name="Dela Cruz",
        email="juan.one@ccash.test",
        phone="09991234568",
    )

    with pytest.raises(ValidationError):
        await service.create_member(
            id_no="000000100",
            first_name="Pedro",
            last_name="Reyes",
            email="pedro.two@ccash.test",
            phone="09991234569",
        )


@pytest.mark.parametrize(
    "id_no,phone",
    [("12345", "09991234567"), ("0000000ab", "09991234567"), ("000000100", "12345")],
)
async def test_admin_create_member_rejects_malformed_id_or_phone(session, id_no, phone):
    from app.domains.admin.service import AdminService

    service = AdminService(session)
    with pytest.raises(ValidationError):
        await service.create_member(
            id_no=id_no,
            first_name="Juan",
            last_name="Dela Cruz",
            email="juan.bad@ccash.test",
            phone=phone,
        )


# --------------------------------------------------- merchant ID resolution


async def test_resolve_merchant_id_no_generates_when_not_supplied(session):
    """No redis needed: _resolve_merchant_id_no never touches it."""
    from app.domains.merchants.service import MerchantService

    service = MerchantService(session, None)
    generated = await service._resolve_merchant_id_no(None)
    assert generated.startswith("M")
    assert is_valid_merchant_id_no(generated)


async def test_resolve_merchant_id_no_rejects_malformed_supplied_value(session):
    from app.domains.merchants.service import MerchantService

    service = MerchantService(session, None)
    with pytest.raises(ValidationError):
        await service._resolve_merchant_id_no("123456789")  # missing the "M" prefix


async def test_resolve_merchant_id_no_rejects_duplicate_supplied_value(session):
    from app.domains.merchants.models import MerchantProfile
    from app.domains.merchants.service import MerchantService

    admin_service_user, _wallet = await _make_bare_user(session, "existing-merchant@ccash.test", "09991234570")
    session.add(
        MerchantProfile(
            user_id=admin_service_user.id,
            merchant_id_no="M111111111",
            company_name="Existing Store",
            contact_person="Owner",
            mobile_no="09991234570",
            address="Somewhere",
            tin="000-000-000",
        )
    )
    await session.commit()

    service = MerchantService(session, None)
    with pytest.raises(ValidationError):
        await service._resolve_merchant_id_no("M111111111")


async def _make_bare_user(session, email: str, phone: str):
    from app.core.security import hash_password
    from app.domains.auth.models import User, UserRole, UserStatus

    user = User(
        email=email,
        phone=phone,
        password_hash=hash_password("Test123!"),
        status=UserStatus.ACTIVE,
        is_verified=True,
        role=UserRole.MERCHANT,
    )
    session.add(user)
    await session.flush()
    return user, None
