"""Transaction history report generation (PDF via reportlab, Excel via openpyxl).

Shared by the self-service export (a Member or Merchant downloading or
emailing their own history) and the admin per-account export. Both consume
the same ``TransactionView`` list already used to render the in-app
Transactions page, so the report matches what the account holder sees.
"""

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.money import format_php
from app.domains.transactions.views import TransactionView

COLUMN_HEADERS = ["Date", "Reference", "Type", "Direction", "Counterparty", "Amount", "Status", "Description"]
BRAND_COLOR = colors.HexColor("#0F6ECD")
STRIPE_COLOR = colors.HexColor("#F5F7FA")
GRID_COLOR = colors.HexColor("#CCCCCC")


def _row_for(view: TransactionView) -> list[str]:
    tx = view.transaction
    counterparty = "-"
    if view.counterparty:
        counterparty = view.counterparty.name or view.counterparty.masked_mobile or "-"
    created = tx.created_at.strftime("%Y-%m-%d %H:%M") if tx.created_at else ""
    signed_amount = tx.amount_cents if view.direction.value == "IN" else -tx.amount_cents
    return [
        created,
        tx.reference or "-",
        tx.type.value,
        view.direction.value,
        counterparty,
        format_php(signed_amount),
        tx.status.value,
        tx.description or "",
    ]


def _generated_at() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def generate_transaction_history_pdf(
    account_label: str, account_meta: list[tuple[str, str]], views: list[TransactionView]
) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.6 * inch, bottomMargin=0.6 * inch)
    styles = getSampleStyleSheet()

    elements = [
        Paragraph("Campe Wallet Transaction History", styles["Title"]),
        Paragraph(account_label, styles["Heading2"]),
    ]
    for label, value in account_meta:
        elements.append(Paragraph(f"<b>{label}:</b> {value}", styles["Normal"]))
    elements.append(Paragraph(f"Generated: {_generated_at()}", styles["Normal"]))
    elements.append(Spacer(1, 0.25 * inch))

    data = [COLUMN_HEADERS] + [_row_for(v) for v in views]
    if len(data) == 1:
        elements.append(Paragraph("No transactions found.", styles["Normal"]))
    else:
        table = Table(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), BRAND_COLOR),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("GRID", (0, 0), (-1, -1), 0.5, GRID_COLOR),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, STRIPE_COLOR]),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )
        )
        elements.append(table)

    doc.build(elements)
    return buffer.getvalue()


def generate_transaction_history_excel(
    account_label: str, account_meta: list[tuple[str, str]], views: list[TransactionView]
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transaction History"

    ws.append(["Campe Wallet Transaction History"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([account_label])
    for label, value in account_meta:
        ws.append([f"{label}:", value])
    ws.append(["Generated:", _generated_at()])
    ws.append([])

    header_row_index = ws.max_row + 1
    ws.append(COLUMN_HEADERS)
    for cell in ws[header_row_index]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F6ECD")

    for view in views:
        ws.append(_row_for(view))

    for i in range(1, len(COLUMN_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
