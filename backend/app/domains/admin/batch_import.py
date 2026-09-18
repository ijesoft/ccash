"""Batch member import: parse an admin-uploaded CSV or XLSX (~3000 rows,
the initial CAMPE membership roll) into row dicts for AdminService.create_member.

Expected columns (header names are matched case-insensitively, any order):
id_no, last_name, first_name, middle_name (optional), mobile, email.
"""

import csv
import io

from openpyxl import load_workbook

REQUIRED_FIELDS = ["id_no", "last_name", "first_name", "mobile", "email"]

_HEADER_ALIASES: dict[str, set[str]] = {
    # Trailing "." is stripped by _normalize_header before matching, so
    # "ID No." and "Mobile No." need no separate dotted entry here.
    "id_no": {"id_no", "id no", "id number", "employee id", "9-digit id no", "9 digit id no", "idno"},
    "last_name": {"last_name", "last name", "lastname", "surname"},
    "first_name": {"first_name", "first name", "firstname", "given name"},
    "middle_name": {"middle_name", "middle name", "middlename"},
    "mobile": {"mobile", "mobile number", "mobile no", "phone", "phone number", "contact number"},
    "email": {"email", "email address"},
}


class BatchImportError(Exception):
    pass


def _normalize_header(raw: str) -> str | None:
    # Strip trailing punctuation (e.g. "ID No.") so "id no." and "id no" match
    # the same alias without having to enumerate every variant.
    key = (raw or "").strip().lower().rstrip(".")
    for field, aliases in _HEADER_ALIASES.items():
        if key in aliases:
            return field
    return None


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _to_dicts(rows: list[list]) -> list[dict]:
    if not rows:
        raise BatchImportError("File is empty")

    field_by_col: dict[int, str] = {}
    for i, raw in enumerate(rows[0]):
        field = _normalize_header(_cell_to_str(raw))
        if field:
            field_by_col[i] = field

    missing = [f for f in REQUIRED_FIELDS if f not in field_by_col.values()]
    if missing:
        raise BatchImportError(f"Missing required column(s): {', '.join(missing)}")

    records = []
    for row in rows[1:]:
        if not any(_cell_to_str(c) for c in row):
            continue  # skip blank rows
        record = {field: "" for field in _HEADER_ALIASES}
        for i, field in field_by_col.items():
            record[field] = _cell_to_str(row[i]) if i < len(row) else ""
        records.append(record)
    return records


def _rows_from_csv(content: bytes) -> list[list]:
    text = content.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


def _rows_from_xlsx(content: bytes) -> list[list]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    return [list(row) for row in ws.iter_rows(values_only=True)]


def parse_member_rows(content: bytes, filename: str) -> list[dict]:
    lower = (filename or "").lower()
    if lower.endswith(".xlsx"):
        rows = _rows_from_xlsx(content)
    elif lower.endswith(".csv"):
        rows = _rows_from_csv(content)
    else:
        raise BatchImportError("Unsupported file type: please upload .csv or .xlsx")
    return _to_dicts(rows)
