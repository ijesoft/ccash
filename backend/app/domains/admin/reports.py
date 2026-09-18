"""Platform-wide transaction ledger export (admin only): PDF/Excel.

Unlike the per-account report (transactions/reports.py), this is not
per-viewer: every row shows both sides of a transfer directly, for
reconciliation across every Member and Merchant account ("for tracking").
"""

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.money import format_php

COLUMN_HEADERS = ["Date", "Reference", "Type", "Status", "From", "To", "Amount", "Fee", "Net Amount", "Description"]
BRAND_COLOR = colors.HexColor("#0F6ECD")
STRIPE_COLOR = colors.HexColor("#F5F7FA")
GRID_COLOR = colors.HexColor("#CCCCCC")


def _row_for(row: dict) -> list[str]:
    return [
        row["created_at"],
        row["reference"] or "-",
        row["type"],
        row["status"],
        row["from"] or "-",
        row["to"] or "-",
        format_php(row["amount_cents"]),
        format_php(row["fee_cents"]),
        format_php(row["net_amount_cents"]),
        row["description"] or "",
    ]


def _generated_at() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def generate_all_transactions_pdf(rows: list[dict]) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), topMargin=0.5 * inch, bottomMargin=0.5 * inch)
    styles = getSampleStyleSheet()

    elements = [
        Paragraph("CCash Platform Transaction Report", styles["Title"]),
        Paragraph(f"Generated: {_generated_at()}", styles["Normal"]),
        Paragraph(f"Total transactions: {len(rows)}", styles["Normal"]),
        Spacer(1, 0.25 * inch),
    ]

    data = [COLUMN_HEADERS] + [_row_for(r) for r in rows]
    if len(data) == 1:
        elements.append(Paragraph("No transactions found.", styles["Normal"]))
    else:
        table = Table(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), BRAND_COLOR),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTSIZE", (0, 0), (-1, -1), 6),
                    ("GRID", (0, 0), (-1, -1), 0.5, GRID_COLOR),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, STRIPE_COLOR]),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )
        )
        elements.append(table)

    doc.build(elements)
    return buffer.getvalue()


def generate_all_transactions_excel(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "All Transactions"

    ws.append(["CCash Platform Transaction Report"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated: {_generated_at()}"])
    ws.append([f"Total transactions: {len(rows)}"])
    ws.append([])

    header_row_index = ws.max_row + 1
    ws.append(COLUMN_HEADERS)
    for cell in ws[header_row_index]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F6ECD")

    for row in rows:
        ws.append(_row_for(row))

    for i in range(1, len(COLUMN_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
