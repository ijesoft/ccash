"""Excel import template for the Master List batch upload.

Header spellings here must stay identical to
``app.domains.admin.batch_import._HEADER_ALIASES`` or the generated file
will be rejected by ``parse_member_rows``.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

TEMPLATE_FILENAME = "master_list_import_template.xlsx"

TEMPLATE_HEADERS = [
    "9-DIGIT ID NO.",
    "Last Name",
    "First Name",
    "Middle Name",
    "Mobile Number",
    "Email Address",
]


def build_template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Master List"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F6F43", end_color="1F6F43", fill_type="solid")
    for col, header in enumerate(TEMPLATE_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    widths = [16, 20, 20, 20, 18, 32]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    ws.freeze_panes = "A2"
    # Data starts on row 2; row 1 is the header. All rows import as ACTIVE.

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
